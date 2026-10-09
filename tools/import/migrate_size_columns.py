"""One-time migration: module_data.xlsx dimensions from height x width to width x height.

Run:
    python tools/import/migrate_size_columns.py
Requires openpyxl. Creates a .bak backup before editing.
"""
from pathlib import Path
import re
import shutil

from openpyxl import load_workbook

SOURCE = Path(__file__).resolve().parents[1] / "data_source" / "module_data.xlsx"
OLD_NAMES = {
    "高X宽": "宽x高", "高x宽": "宽x高", "高×宽": "宽x高",
    "炮塔高X宽": "炮塔宽x高", "炮塔高x宽": "炮塔宽x高", "炮塔高×宽": "炮塔宽x高",
}
SIZE_RE = re.compile(r"^\s*(\d+)\s*[xX×]\s*(\d+)\s*$")


def migrate(path: Path = SOURCE) -> None:
    book = load_workbook(path)
    changed = 0
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or cell.value not in OLD_NAMES:
                    continue
                for cells in sheet.iter_cols(min_col=cell.column, max_col=cell.column, min_row=cell.row + 1):
                    for value_cell in cells:
                        if value_cell.value is None or str(value_cell.value).strip() == "":
                            continue
                        found = SIZE_RE.fullmatch(str(value_cell.value))
                        if not found:
                            raise ValueError(f"{sheet.title}!{value_cell.coordinate}: invalid size {value_cell.value!r}")
                        value_cell.value = f"{found[2]}x{found[1]}"
                cell.value = OLD_NAMES[cell.value]
                changed += 1
    if not changed:
        print("No old size headings; workbook left unchanged.")
        return
    backup = path.with_suffix(".xlsx.bak")
    shutil.copy2(path, backup)
    book.save(path)
    print(f"Converted {changed} headers and their values; backup: {backup}")


if __name__ == "__main__":
    migrate()
